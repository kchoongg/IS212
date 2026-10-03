#!/usr/bin/env python3
"""
Fill the team's Excel test case template from a JSON file.

Usage:
    python fill_template.py <test_cases.json> <output.xlsx> [--legacy]

Checks the JSON first, including the team's writing rules (every expected
result tied to a "Step N:", no hand-numbered, filler or "Repeat" steps, case
references that exist, dependent cases placed right after what they depend
on, no full reseed as a reset). --legacy downgrades the writing rules to a
warning, for rebuilding a file written before 2026-09-30 without changing it.

The layout of each test case sheet comes from Test_Case_Template.xlsx (sheet "TC01").
This script copies that sheet once per test case, so the format never drifts.
It also adds a Summary sheet (traceability + coverage) and, when any question is
still unanswered, an Open Questions sheet. Entries starting "RESOLVED" stay in
the JSON as the decision record and are left out of the workbook.

Requires: pip install openpyxl
"""
import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    sys.exit("openpyxl is not installed. Run: pip install openpyxl")

TEMPLATE = Path(__file__).resolve().parent.parent / "Test_Case_Template.xlsx"
TEMPLATE_SHEET = "TC01"

# Cells in the template (column B holds the values)
FIELD_CELLS = {
    "id": "B2",
    "preconditions": "B3",
    "steps": "B4",
    "test_data": "B5",
    "expected_result": "B6",
    "created_by": "B7",
    "creation_date": "B8",
}
REQUIRED_TC_FIELDS = ["id", "title", "linked_ac", "type", "source",
                      "preconditions", "steps", "test_data", "expected_result"]
VALID_TYPES = {"Positive", "Negative", "Boundary"}
VALID_SOURCES = {"Acceptance Criteria", "Code"}

ARIAL = "Arial"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="FFA4C2F4")  # same blue as template header
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def validate(data):
    """Return a list of problems. Empty list means the JSON is good."""
    problems = []
    for key in ["story_key", "story_title", "created_by", "acceptance_criteria", "test_cases"]:
        if not data.get(key):
            problems.append(f"Missing top-level field: {key}")
    if problems:
        return problems

    ac_ids = [ac.get("id") for ac in data["acceptance_criteria"]]
    seen_tc = set()
    covered = set()
    for i, tc in enumerate(data["test_cases"], start=1):
        label = tc.get("id") or f"test case #{i}"
        for f in REQUIRED_TC_FIELDS:
            if not tc.get(f):
                problems.append(f"{label}: missing '{f}'")
        if tc.get("id") in seen_tc:
            problems.append(f"{label}: duplicate test case id")
        seen_tc.add(tc.get("id"))
        if tc.get("type") and tc["type"] not in VALID_TYPES:
            problems.append(f"{label}: type must be one of {sorted(VALID_TYPES)}")
        if tc.get("source") and tc["source"] not in VALID_SOURCES:
            problems.append(f"{label}: source must be one of {sorted(VALID_SOURCES)}")
        for ac in tc.get("linked_ac", []):
            if ac not in ac_ids:
                problems.append(f"{label}: linked_ac '{ac}' is not in acceptance_criteria")
            covered.add(ac)
    for ac in ac_ids:
        if ac not in covered:
            problems.append(f"{ac} has no test case covering it")
    return problems


# The team's rules for how a test case is written (SKILL.md, Step 3). Checked
# here so a case that breaks one cannot be built, whoever generated it.
STEP_LINE = re.compile(r"^Step (\d+): \S")
PRE_NUMBERED = re.compile(r"^\s*\d+[.)]\s")
FILLER_STEP = re.compile(r"^(Read|Look at|Review|Check) the (list|page|screen|results?)\b", re.I)
REPEAT_STEP = re.compile(r"^Repeat\b", re.I)
# A full reseed wipes every table on the shared database. Mentioning it is fine
# ("do NOT run `npx prisma db seed`"); telling the tester to run it is not.
RESEED = re.compile(r"(?<!not )(?<!NOT )(?<!never )\brun `?npx prisma db seed`?", re.I)


def check_writing_rules(data):
    problems = []
    key = data["story_key"]
    tcs = data["test_cases"]
    ids = [tc.get("id") for tc in tcs]
    ref_pattern = re.compile(rf"\b{re.escape(key)}-TC\d+\b")

    for i, tc in enumerate(tcs):
        label = tc.get("id") or f"test case #{i + 1}"
        steps = tc.get("steps") or []
        if not isinstance(steps, list):
            problems.append(f"{label}: 'steps' must be a list (the script numbers them)")
            steps = []

        for n, step in enumerate(steps, start=1):
            if PRE_NUMBERED.match(step):
                problems.append(f"{label} step {n}: don't number steps yourself - the script does")
            if FILLER_STEP.match(step):
                problems.append(f"{label} step {n}: '{step[:40]}' is a filler step - put its result on the click that shows it")
            if REPEAT_STEP.match(step):
                problems.append(f"{label} step {n}: write the repeated steps out, so each has its own result")

        # Every expected result is tied to the step that produces it.
        last = 0
        for line in str(tc.get("expected_result", "")).split("\n"):
            if not line.strip():
                continue
            m = STEP_LINE.match(line)
            if not m:
                problems.append(f"{label}: expected result line must start 'Step N: ' - got '{line[:50]}'")
                continue
            n = int(m.group(1))
            if not 1 <= n <= len(steps):
                problems.append(f"{label}: 'Step {n}' but the case has {len(steps)} step(s)")
            elif n <= last:
                problems.append(f"{label}: 'Step {n}' is out of order or repeated - one line per step, in step order")
            last = max(last, n)

        text = " ".join([str(tc.get("preconditions", "")), str(tc.get("expected_result", "")),
                         str(tc.get("test_data", ""))] + [str(s) for s in steps])
        for ref in ref_pattern.findall(text):
            if ref not in ids:
                problems.append(f"{label}: refers to {ref}, which doesn't exist")

        # A case that needs another run first comes straight after it.
        for m in re.finditer(rf"({re.escape(key)}-TC\d+) has just been run", str(tc.get("preconditions", ""))):
            if i == 0 or ids[i - 1] != m.group(1):
                problems.append(f"{label}: depends on {m.group(1)}, so it must come immediately after it")

        if RESEED.search(str(tc.get("preconditions", ""))):
            problems.append(f"{label}: don't reset data with a full reseed - use the story's fixture script "
                            "(e.g. `npm run fixtures:spm28`), which only touches its own test events")
    return problems


def row_height(text, chars_per_line=60, line_pt=13.5):
    """Estimate a row height so wrapped text is visible when opened in Excel."""
    lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))
    return max(15, lines * line_pt + 4)


def as_text(value):
    if isinstance(value, list):
        return "\n".join(str(v) for v in value)
    return str(value)


def write_cell(ws, ref, value, bold=False, fill=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=ARIAL, size=10, bold=bold)
    c.border = BORDER
    c.alignment = WRAP_TOP
    if fill:
        c.fill = fill
    return c


def still_open(data):
    """The open_questions entries that are not marked RESOLVED."""
    return [q for q in data.get("open_questions") or []
            if not as_text(q).lstrip().upper().startswith("RESOLVED")]


def build(data, out_path):
    wb = load_workbook(TEMPLATE)
    if TEMPLATE_SHEET not in wb.sheetnames:
        sys.exit(f"Template sheet '{TEMPLATE_SHEET}' not found in {TEMPLATE}")
    base = wb[TEMPLATE_SHEET]
    base.title = "_template"

    created = data.get("creation_date") or date.today().isoformat()
    created_date = datetime.strptime(created, "%Y-%m-%d").date()

    sheet_names = []
    for n, tc in enumerate(data["test_cases"], start=1):
        ws = wb.copy_worksheet(base)
        ws.title = f"TC{n:02d}"
        sheet_names.append(ws.title)

        steps = tc["steps"]
        if isinstance(steps, list):
            steps = "\n".join(f"{i}. {s}" for i, s in enumerate(steps, start=1))

        values = {
            "id": tc["id"],
            "preconditions": as_text(tc["preconditions"]),
            "steps": steps,
            "test_data": as_text(tc["test_data"]),
            "expected_result": as_text(tc["expected_result"]),
            "created_by": data["created_by"],
            "creation_date": created_date,
        }
        for field, ref in FIELD_CELLS.items():
            cell = ws[ref]
            cell.value = values[field]
            cell.alignment = WRAP_TOP
            ws[ref.replace("B", "A")].alignment = Alignment(vertical="top")
            if field == "creation_date":
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="left", vertical="top")
            else:
                ws.row_dimensions[cell.row].height = row_height(cell.value)

    wb.remove(base)

    # ---- Summary sheet (traceability) ----
    summary = wb.create_sheet("Summary", 0)
    write_cell(summary, "A1", "Story", bold=True, fill=HEADER_FILL)
    write_cell(summary, "B1", f'{data["story_key"]}: {data["story_title"]}', bold=True, fill=HEADER_FILL)
    summary.merge_cells("B1:F1")
    summary.row_dimensions[1].height = row_height(summary["B1"].value, 110)

    headers = ["Sheet", "Test Case ID", "Title", "Linked AC", "Type", "Source"]
    for col, h in enumerate(headers, start=1):
        write_cell(summary, f"{'ABCDEF'[col-1]}3", h, bold=True, fill=HEADER_FILL)
    for r, (tc, sheet) in enumerate(zip(data["test_cases"], sheet_names), start=4):
        link = write_cell(summary, f"A{r}", sheet)
        link.hyperlink = f"#'{sheet}'!A1"
        link.font = Font(name=ARIAL, size=10, color="FF1155CC", underline="single")
        write_cell(summary, f"B{r}", tc["id"])
        write_cell(summary, f"C{r}", tc["title"])
        write_cell(summary, f"D{r}", ", ".join(tc["linked_ac"]))
        write_cell(summary, f"E{r}", tc["type"])
        write_cell(summary, f"F{r}", tc["source"])

    r = len(data["test_cases"]) + 6
    for ref, h in [("A", "AC"), ("B", "Acceptance Criterion"), ("E", "Covered by")]:
        write_cell(summary, f"{ref}{r}", h, bold=True, fill=HEADER_FILL)
    summary.merge_cells(f"B{r}:D{r}")
    summary.merge_cells(f"E{r}:F{r}")
    for ac in data["acceptance_criteria"]:
        r += 1
        ids = [tc["id"] for tc in data["test_cases"] if ac["id"] in tc["linked_ac"]]
        write_cell(summary, f"A{r}", ac["id"])
        write_cell(summary, f"B{r}", ac["text"])
        summary.merge_cells(f"B{r}:D{r}")
        write_cell(summary, f"E{r}", "\n".join(ids))
        summary.merge_cells(f"E{r}:F{r}")
        summary.row_dimensions[r].height = max(row_height(ac["text"], 70), 13.5 * len(ids) + 4)

    if data.get("code_references"):
        r += 2
        write_cell(summary, f"A{r}", "Code reviewed", bold=True, fill=HEADER_FILL)
        write_cell(summary, f"B{r}", "\n".join(data["code_references"]))
        summary.merge_cells(f"B{r}:F{r}")
        summary.row_dimensions[r].height = 13.5 * len(data["code_references"]) + 4

    for col, width in zip("ABCDEF", [12, 20, 48, 12, 18, 20]):
        summary.column_dimensions[col].width = width

    # ---- Open Questions sheet ----
    # Only questions still waiting for an answer. A "RESOLVED <date>: ..." entry
    # stays in the JSON as the record of the decision but is not a question for
    # the reader, and a workbook with nothing left to ask gets no sheet at all.
    questions = still_open(data)
    if questions:
        oq = wb.create_sheet("Open Questions", 1)
        write_cell(oq, "A1", "#", bold=True, fill=HEADER_FILL)
        write_cell(oq, "B1", "Question / ambiguity for the PO or developer", bold=True, fill=HEADER_FILL)
        write_cell(oq, "C1", "Answer", bold=True, fill=HEADER_FILL)
        for i, q in enumerate(questions, start=2):
            write_cell(oq, f"A{i}", i - 1)
            write_cell(oq, f"B{i}", q)
            write_cell(oq, f"C{i}", None)
            oq.row_dimensions[i].height = row_height(q, 70)
        oq.column_dimensions["A"].width = 5
        oq.column_dimensions["B"].width = 75
        oq.column_dimensions["C"].width = 40

    wb.active = 0
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    # Test cases quote on-screen text ("−", "—", "…"); a Windows console's
    # default code page can't print them, which would crash the report.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    legacy = "--legacy" in sys.argv[1:]
    if len(args) != 2:
        sys.exit(__doc__)
    json_path, out_path = Path(args[0]), Path(args[1])
    data = json.loads(json_path.read_text(encoding="utf-8"))

    problems = validate(data)
    if problems:
        print("VALIDATION FAILED - fix the JSON and run again:")
        for p in problems:
            print(f"  - {p}")
        sys.exit(1)

    # The writing rules (SKILL.md Step 3) are enforced for every new or revised
    # file. --legacy only exists so test cases written before 2026-09-30 can
    # still be rebuilt unchanged; it reports the problems instead of refusing.
    rule_problems = check_writing_rules(data)
    if rule_problems and not legacy:
        print("VALIDATION FAILED - the test cases break the writing rules in SKILL.md Step 3:")
        for p in rule_problems:
            print(f"  - {p}")
        print("(Rebuilding an unchanged file written before 2026-09-30? Add --legacy.)")
        sys.exit(1)
    if rule_problems:
        print(f"WARNING (--legacy): {len(rule_problems)} writing-rule problem(s) left in - "
              "fix them next time these test cases are revised.")

    build(data, out_path)
    print(f"OK: wrote {len(data['test_cases'])} test cases for "
          f"{len(data['acceptance_criteria'])} acceptance criteria to {out_path}")


if __name__ == "__main__":
    main()
